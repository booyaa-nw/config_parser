# core/excel/style.py
#
# 旧 booya_config_parser/export/style.py とほぼ同内容。セルの見た目
# (罫線・ヘッダ塗りつぶし・折り返し)はレイアウトの種類に依存しないため、
# layouts.py の各レンダラーから共通で呼ばれる。
from __future__ import annotations

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.borders import Border, Side
from openpyxl.utils import get_column_letter

table_header_fill = PatternFill(patternType="solid", fgColor="d3d3d3")
table_header_font = Font(bold=True)
table_side = Side(style="thin", color="000000")
table_border = Border(top=table_side, bottom=table_side, left=table_side, right=table_side)
cell_alignment = Alignment(wrapText=True)

sheet_title_font = Font(bold=True, size=16)
subtable_title_font = Font(bold=True, size=11)


def write_sheet_title(ws, row, col, value):
    ws.cell(row=row, column=col).value = value
    ws.cell(row=row, column=col).font = sheet_title_font
    return ws


def write_subtable_title(ws, row, col, value):
    """compositeレイアウトで、メインブロックの下に続くサブテーブルの見出し。"""
    ws.cell(row=row, column=col).value = value
    ws.cell(row=row, column=col).font = subtable_title_font
    return ws


def write_header_cell(ws, row, col, value, width=None):
    ws.cell(row=row, column=col).value = value
    ws.cell(row=row, column=col).font = table_header_font
    ws.cell(row=row, column=col).border = table_border
    ws.cell(row=row, column=col).fill = table_header_fill
    if width:
        ws.column_dimensions[get_column_letter(col)].width = width
    return ws


def write_cell(ws, row, col, value, width=None):
    ws.cell(row=row, column=col).value = value
    ws.cell(row=row, column=col).alignment = cell_alignment
    ws.cell(row=row, column=col).border = table_border
    if width:
        ws.column_dimensions[get_column_letter(col)].width = width
    return ws


def cell_text(value) -> str:
    """フィールド値(常にlist、まれにstr/その他)を1セルのテキストへ変換する共通ヘルパー。
    複数値は改行区切りでセル内改行表示する(write_cellのwrapText=Trueと組み合わせ)。
    """
    if isinstance(value, list):
        return "\n".join(str(v) for v in value)
    if value is None:
        return ""
    return str(value)
