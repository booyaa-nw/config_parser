# core/excel/writer.py
#
# 旧 export/excel.py の ExportExcel クラスの置き換え。
#
# 旧実装との最大の違い: export_book()は「あらかじめ全カテゴリ分
# 整形済みのformatted_config_obj」を丸ごと受け取り、category名を
# 直接見て分岐していた。新設計ではCategoryFormatterのリスト
# (=実装済みカテゴリの一覧、1ファイル=1シートの原則そのもの)を受け取り、
# 各Formatterに対してdoc全体を渡して呼び出す。
#
# これにより:
#   - 未実装カテゴリ(旧実装の `elif category == 'ipsec_vpn': continue`)は
#     formattersのリストに含めないだけで自動的にスキップされる、
#     export層側の特別扱いが不要になった。
#   - system_ha専用だったexport/ha.pyは、field_names側の
#     `{'subtable': {...}}` 宣言を読む汎用compositeレンダラーに
#     置き換わり(core/excel/layouts.py)、export層はsystem_haという
#     カテゴリ名を一切知らない。
#
# **2026-10-22追記(vpn_ipsec_phase1/phase2のExcelレイアウト例外対応)**:
# こうぢさんの依頼で、vpn_ipsec_phase1/phase2だけはExcel出力に限り
# 1シートに統合する例外レイアウトに対応した(JSON出力側のカテゴリ分割は
# 変更なし)。`field_names['__layout__']['excel_sheet_group']`が同じ値を
# 持つ、FORMATTERSリスト上で連続したFormatter群を1シートにまとめる
# (`_group_formatters_by_sheetname()`)。これもcategory名のハードコード
# ではなく、あくまでField_names側の宣言(どのFormatterがどれと組むか)で
# 決まる、という既存方針のまま実現している。
from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path

from openpyxl import Workbook

from ..doc_access import iter_vdom_names
from ..naming import make_export_base_name
from . import style
from .font import set_japanese_font
from .layouts import render, render_paired_vertical_group

_INVALID_SHEETNAME_CHARS = re.compile(r"[:\\/?*\[\]]")
_MAX_SHEETNAME_LEN = 31


class ExcelWriter:
    """formatters: CategoryFormatterインスタンスのリスト
    (通常は ftnt/fos/formatters/__init__.py の FORMATTERS のような、
    vendor/機種ファミリーごとの登録済み一覧をそのまま渡す)。
    """

    def __init__(self, formatters: list):
        self.formatters = formatters

    def export(self, doc: dict, export_dir: str = "./tmp/export", base_name: str | None = None) -> Path:
        """base_name: 未指定ならmake_export_base_name(doc)で生成する。
        core/pipeline.pyがExcel/JSON両出力で同じbase_nameを使い回すために
        指定できるようにしてある(2026-09-26追加、CLI実装に伴う変更)。"""
        export_dir_path = Path(export_dir)
        export_dir_path.mkdir(parents=True, exist_ok=True)
        is_multivdom = doc["header"]["vdom"] == 1

        wb = Workbook()
        wb.remove(wb.worksheets[0])
        used_sheetnames: set[str] = set()

        global_formatters = [f for f in self.formatters if f.resolve_field_names(doc)["__layout__"].get("global", False)]
        vdom_formatters = [f for f in self.formatters if f not in global_formatters]

        # excel_sheet_groupが同じ値を持つ、リスト上で連続したFormatter群は
        # 1シートに統合する(vpn_ipsec_phase1/phase2用の例外対応、上記
        # ヘッダコメント参照)。グルーピングはdoc依存のresolve_field_names()
        # 結果だけで決まり、vdomには依存しないため、vdomループの外(1回)で
        # 計算しておく。
        global_groups = self._group_formatters_by_sheetname(global_formatters, doc)
        vdom_groups = self._group_formatters_by_sheetname(vdom_formatters, doc)

        # global系(system_global/system_ha等、vdomに属さないカテゴリ)を先頭に
        # まとめて出す。vdom単位ではないので並び替えの対象外。
        for group in global_groups:
            self._write_group_sheet(wb, doc, group, "global", used_sheetnames, is_multivdom)

        # 2026-09-26変更: シート順序をFormatter(カテゴリ)単位ではなく
        # vdom単位でグルーピングするよう変更(こうぢさんの依頼: "root address,
        # root address group, root ..., ext-vdom address, ext-vdom address
        # group, ext-vdom ..."のように、vdomごとにシートをまとめたい)。
        # 変更前はFormatter順を外側・vdomを内側にしていたため、
        # "root address, ext-vdom address, root address group, ..."のように
        # カテゴリ単位でインターリーブされていた。
        for vdom_name in iter_vdom_names(doc):
            for group in vdom_groups:
                self._write_group_sheet(wb, doc, group, vdom_name, used_sheetnames, is_multivdom)

        if base_name is None:
            base_name = make_export_base_name(doc)
        book_path = self._unique_book_path(export_dir_path, base_name)
        wb.save(book_path)
        set_japanese_font(book_path)
        return book_path

    @staticmethod
    def _group_formatters_by_sheetname(formatters: list, doc: dict) -> list[list]:
        """`__layout__.excel_sheet_group`(未指定なら`sheetname`)が同じ値を
        持つ、リスト上で連続したFormatterを1グループにまとめる。
        通常のカテゴリはexcel_sheet_groupを宣言しないため、各自の
        sheetnameがそのままグループキーになり、結果的に1Formatter=1グループ
        (=従来どおり1シート)になる。vpn_ipsec_phase1/phase2のような
        意図的な例外だけが2件以上のグループを持つ。"""
        groups: list[list] = []
        current: list = []
        current_key = None
        for formatter in formatters:
            layout = formatter.resolve_field_names(doc)["__layout__"]
            key = layout.get("excel_sheet_group", layout["sheetname"])
            if current and key == current_key:
                current.append(formatter)
            else:
                if current:
                    groups.append(current)
                current = [formatter]
                current_key = key
        if current:
            groups.append(current)
        return groups

    @staticmethod
    def _write_group_sheet(wb, doc, group: list, vdom_name, used_sheetnames, is_multivdom) -> None:
        if len(group) == 1:
            ExcelWriter._write_formatter_sheet(wb, doc, group[0], vdom_name, used_sheetnames, is_multivdom)
        else:
            ExcelWriter._write_paired_vertical_sheet(wb, doc, group, vdom_name, used_sheetnames, is_multivdom)

    @staticmethod
    def _write_paired_vertical_sheet(wb, doc, group: list, vdom_name, used_sheetnames, is_multivdom) -> None:
        """excel_sheet_groupでまとめられたFormatter群(先頭=親, 以降=子)を
        1シートにまとめて書き出す(vpn_ipsec_phase1/phase2専用の例外対応、
        本ファイル先頭のヘッダコメント参照)。親カテゴリのレコードが0件なら
        (=そのvdomにphase1トンネル自体が無い)、子だけ表示する意味が無い
        ため、既存の「空カテゴリはシートを作らない」方針のままシート自体を
        作らない。"""
        parts = [formatter.format(doc, vdom_name) for formatter in group]
        primary_records, primary_field_names = parts[0]
        if not primary_records:
            return

        primary_layout = primary_field_names["__layout__"]
        is_global = primary_layout.get("global", False)
        base_sheetname = primary_layout["excel_sheet_group"]
        if is_global or not is_multivdom:
            label = base_sheetname
        else:
            label = f"{vdom_name} {base_sheetname}"
        sheetname = ExcelWriter._unique_sheetname(label, used_sheetnames)

        ws = wb.create_sheet(sheetname)
        style.write_sheet_title(ws, 1, 1, sheetname)

        # 各パートのセクション見出し(sheetname)にも、シート名と同じ規則で
        # vdom名を前置きする(renderingは見た目のラベルしか見ないので、
        # ここでvdom名を合成してから渡す)。
        labeled_parts = []
        for records, field_names in parts:
            field_names = deepcopy(field_names)
            base_title = field_names["__layout__"]["sheetname"]
            if is_global or not is_multivdom:
                title = base_title
            else:
                title = f"{vdom_name} {base_title}"
            field_names["__layout__"]["sheetname"] = title
            labeled_parts.append((records, field_names))

        # header_row=2: シート全体の見出し(row1)の直後が親セクション自身の
        # 見出し行になる(シート見出しと親セクション見出しの間は空けない。
        # サンプル画像`ipsec_vpn_image.xlsx`の行1/行2の並びに合わせている)。
        render_paired_vertical_group(ws, header_row=2, parts=labeled_parts)

    @staticmethod
    def _write_formatter_sheet(wb, doc, formatter, vdom_name, used_sheetnames, is_multivdom) -> None:
        records, field_names = formatter.format(doc, vdom_name)
        if not records:
            # 空カテゴリはシートを作らない。旧実装は全カテゴリ分無条件で
            # シートを作っていたが、実データだと未使用vdomで大量の空シートが
            # 発生するため変更した(要相談としてarchitecture/config-parser-
            # design.mdに明記)。
            return

        is_global = field_names["__layout__"].get("global", False)
        base_sheetname = field_names["__layout__"]["sheetname"]
        if is_global or not is_multivdom:
            label = base_sheetname
        else:
            label = f"{vdom_name} {base_sheetname}"
        sheetname = ExcelWriter._unique_sheetname(label, used_sheetnames)

        ws = wb.create_sheet(sheetname)
        style.write_sheet_title(ws, 1, 1, sheetname)
        render(ws, header_row=3, records=records, field_names=field_names)

    @staticmethod
    def _unique_book_path(export_dir: Path, base_name: str) -> Path:
        candidate = export_dir / f"{base_name}.xlsx"
        n = 2
        while candidate.exists():
            candidate = export_dir / f"{base_name}_{n}.xlsx"
            n += 1
        return candidate

    @staticmethod
    def _unique_sheetname(label: str, used: set[str]) -> str:
        """Excelのシート名制約(31文字以内・`:\\/?*[]`禁止)への対応。
        旧実装では `# 31文字オーバーの対処 @todo` のまま未対応だった箇所。
        31文字超過時は末尾を削り、切り詰めにより名前が衝突する場合は
        連番を付けて一意にする。"""
        name = _INVALID_SHEETNAME_CHARS.sub("_", label)[:_MAX_SHEETNAME_LEN]
        candidate = name
        n = 2
        while candidate in used:
            suffix = f"_{n}"
            candidate = name[: _MAX_SHEETNAME_LEN - len(suffix)] + suffix
            n += 1
        used.add(candidate)
        return candidate
