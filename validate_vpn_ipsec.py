"""vpn_ipsec_phase1 / vpn_ipsec_phase2 Formatterの実データ検証スクリプト。

実サンプル2件(FG_7.6_config01.conf / FG_7.6_config01_vdom.conf、どちらも
VPNウィザードで作成したdialup接続2件: test_ipsec1, test_remote01)に対して
実際にprocess_file()を動かし、以下を検証する。

- FG_7.2系サンプル(ipsec vpn設定が無い)では ipsec_phase1/ipsec_phase2 の
  カテゴリ自体がJSON出力に現れない(空カテゴリのシート/キーを作らない仕様)。
- FG_7.6系サンプルでは phase1/phase2 ともに2件ずつ、キー集合が
  field_namesの表示列と完全一致する。
- phase2の`src`/`dst`派生ロジック: addr-type=nameのエントリは
  `src-name`/`dst-name`の値がそのまま出ること、addr-type省略(→subnet
  デフォルト)のエントリは `src-subnet`/`dst-subnet`のデフォルト値から
  "0.0.0.0/0.0.0.0" が組み立てられること。
- phase2の`phase1name`で対応するphase1トンネル名と突合せできること。
- 【2026-10-22追記】Excel出力側の例外レイアウト(こうぢさん依頼、
  `ipsec_vpn_image.xlsx`準拠): phase1/phase2が1シート("<vdom> ipsec_vpn")
  に統合され、phase1トンネルが列方向に並び、対応するphase2がその下に
  積まれること(JSON側のカテゴリ分割は従来どおり変更無し)。

実行方法: `uv run --project tools/config_parser python tools/config_parser/validate_vpn_ipsec.py`
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).parent
sys.path.insert(0, str(ROOT.parent / "common" / "src"))
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook  # noqa: E402

from config_parser.ftnt.fos import process_file  # noqa: E402

SAMPLE_DIR = ROOT / "tests" / "fg_config_sample"

NO_IPSEC_SAMPLES = [
    SAMPLE_DIR / "FG_7.2_config01.conf",
    SAMPLE_DIR / "FG_7.2_config01_vdom.conf",
]
IPSEC_SAMPLES = [
    SAMPLE_DIR / "FG_7.6_config01.conf",
    SAMPLE_DIR / "FG_7.6_config01_vdom.conf",
]

EXPECTED_PHASE1_KEYS = {
    "name", "type", "interface", "remote-gw", "ike-version", "peertype",
    "proposal", "dhgrp", "net-device", "mode-cfg", "eap", "transport",
    "ipv4-start-ip", "ipv4-end-ip", "dns-mode", "ipv4-split-include",
    "save-password", "wizard-type", "comments",
}
EXPECTED_PHASE2_KEYS = {
    "name", "phase1name", "proposal", "dhgrp", "src-addr-type", "src",
    "dst-addr-type", "dst", "comments",
}


def load_categories(sample_path, export_dir):
    import json

    result = process_file(str(sample_path), export_dir=str(export_dir), json_only=True)
    merged = {}
    for p in result.json_paths:
        doc = json.loads(p.read_text(encoding="utf-8"))
        for cat, body in doc.get("categories", {}).items():
            merged.setdefault(cat, []).extend(body["records"])
    return merged


def main():
    import tempfile

    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)

        # 1. ipsec vpn設定が無いサンプルでは、カテゴリ自体が出現しないこと
        for sample in NO_IPSEC_SAMPLES:
            cats = load_categories(sample, tmp)
            assert "ipsec_phase1" not in cats, sample
            assert "ipsec_phase2" not in cats, sample
        print("OK: ipsec vpn未設定サンプル(FG_7.2系2件)でカテゴリが出現しないこと")

        # 2. ipsec vpn設定があるサンプルでの検証
        for sample in IPSEC_SAMPLES:
            cats = load_categories(sample, tmp)

            phase1_records = cats["ipsec_phase1"]
            phase2_records = cats["ipsec_phase2"]
            assert len(phase1_records) == 2, (sample, phase1_records)
            assert len(phase2_records) == 2, (sample, phase2_records)

            for r in phase1_records:
                assert set(r) == EXPECTED_PHASE1_KEYS, (sample, set(r))
            for r in phase2_records:
                assert set(r) == EXPECTED_PHASE2_KEYS, (sample, set(r))

            by_name_p1 = {r["name"][0]: r for r in phase1_records}
            by_name_p2 = {r["name"][0]: r for r in phase2_records}

            assert set(by_name_p1) == {"test_ipsec1", "test_remote01"}, sample
            assert set(by_name_p2) == {"test_ipsec1", "test_remote01"}, sample

            # phase1name でphase1側トンネル名と突合せできること
            for name, p2 in by_name_p2.items():
                assert p2["phase1name"] == [name], (sample, name, p2["phase1name"])
                assert name in by_name_p1, (sample, name)

            # test_remote01: mode-cfg/eap/ipv4-start-ip等が明示設定されている側
            remote01 = by_name_p1["test_remote01"]
            assert remote01["mode-cfg"] == ["enable"]
            assert remote01["eap"] == ["enable"]
            assert remote01["ipv4-start-ip"] == ["192.168.222.0"]
            assert remote01["wizard-type"] == ["dialup-forticlient"]

            # test_ipsec1: 上記が未設定側(デフォルト値にフォールバック)
            ipsec1 = by_name_p1["test_ipsec1"]
            assert ipsec1["mode-cfg"] == ["disable"]
            assert ipsec1["eap"] == ["disable"]
            assert ipsec1["ipv4-start-ip"] == []
            assert ipsec1["wizard-type"] == ["dialup-fortigate"]

            # src/dst派生: addr-type=nameのエントリ(test_ipsec1)は-name値、
            # addr-type省略=subnetデフォルトのエントリ(test_remote01)は
            # src-subnet/dst-subnetのデフォルト値から組み立てられること
            p2_ipsec1 = by_name_p2["test_ipsec1"]
            assert p2_ipsec1["src-addr-type"] == ["name"]
            assert p2_ipsec1["src"][0].endswith("_local") or p2_ipsec1["src"][0].endswith("_local02")
            assert p2_ipsec1["dst-addr-type"] == ["name"]
            assert p2_ipsec1["dst"] == ["test_ipsec1_remote"]

            p2_remote01 = by_name_p2["test_remote01"]
            assert p2_remote01["src-addr-type"] == ["subnet"], (sample, p2_remote01)
            assert p2_remote01["src"] == ["0.0.0.0/0.0.0.0"], (sample, p2_remote01)
            assert p2_remote01["dst-addr-type"] == ["subnet"]
            assert p2_remote01["dst"] == ["0.0.0.0/0.0.0.0"]

        print("OK: ipsec vpn設定ありサンプル(FG_7.6系2件)でphase1/phase2ともに"
              "キー集合一致・phase1name突合せ・src/dst派生ロジックを確認")

        # 3. Excel出力側の例外レイアウト(こうぢさん依頼)の確認。
        #    phase1/phase2が別シートにならず、1シート("<vdom> ipsec_vpn")に
        #    統合されていること、phase1トンネルが列方向(A/B, C/D, ...)に
        #    並び、対応するphase2がその下に積まれていることを確認する。
        for sample in IPSEC_SAMPLES:
            result = process_file(str(sample), export_dir=str(tmp))
            wb = load_workbook(result.excel_path)

            # novdomは単に"ipsec_vpn"、multivdomは"root ipsec_vpn"になる
            # (他カテゴリのシート名規則と同じ、vdom名はmultivdomの時だけ
            # 前置きされる)。どちらにせよ"ipsec"を含むシートは1つだけ。
            ipsec_sheets = [name for name in wb.sheetnames if "ipsec" in name]
            assert len(ipsec_sheets) == 1, (sample, wb.sheetnames)
            ws = wb[ipsec_sheets[0]]

            phase1_title_rows = [
                r for r in range(1, ws.max_row + 1)
                if ws.cell(row=r, column=1).value and "ipsec_phase1" in str(ws.cell(row=r, column=1).value)
            ]
            assert len(phase1_title_rows) == 1, (sample, phase1_title_rows)
            phase1_row = phase1_title_rows[0]
            # phase1は列方向(A/B=test_ipsec1, C/D=test_remote01の
            # いずれかの順)に並び、見出し行の2行下(見出し→1行空け→データ)の
            # 先頭フィールドは"name"であること。
            data_row = phase1_row + 2
            assert ws.cell(row=data_row, column=1).value == "name"
            assert ws.cell(row=data_row, column=3).value == "name"
            assert {ws.cell(row=data_row, column=2).value, ws.cell(row=data_row, column=4).value} \
                == {"test_ipsec1", "test_remote01"}

            # phase2の見出し行(値は実データに依存するため、行番号自体は
            # 固定せず見出しがどこかに1箇所だけあることを確認する)。
            phase2_title_rows = [
                r for r in range(1, ws.max_row + 1)
                if ws.cell(row=r, column=1).value and "ipsec_phase2" in str(ws.cell(row=r, column=1).value)
            ]
            assert len(phase2_title_rows) == 1, (sample, phase2_title_rows)

        print("OK: Excel出力がphase1/phase2を1シート('<vdom> ipsec_vpn')に"
              "統合し、phase1を列方向に並べていることを確認")

    print("ALL CHECKS PASSED")


if __name__ == "__main__":
    main()
